"""Welle 4: Prüfregeln, Versandarten, Excel-Übersicht/Summenblatt,
Importsperre über Rechner, Diagnose als Funktion."""
import json
import logging
import os
import time

import openpyxl
import pytest

import diagnose as d
import woo_to_cdh as w
from conftest import FakeWoo


def _abrufen(u, **kw):
    return w.abrufen(u["cfg"], client_factory=FakeWoo, **kw)


def _wex(einheit, tmp_path):
    p = tmp_path / "vorschau.wex"
    w.write_cdh_wex(p, einheit.wex_data)
    return p.read_text(encoding="utf-8-sig")


def _shop(u, name):
    return next(s for s in u["cfg"]["shops"] if s["name"] == name)


# --- Prüfregel: Kundenadresse fehlt → CDH-Standardadresse ------------------

def test_kundenadresse_fehlt_sender_leer(umgebung, tmp_path):
    del _shop(umgebung, "Mitarbeiter-Shop")["sender_address"]
    pruef = _abrufen(umgebung)
    sh = next(s for s in pruef.shops if s.shop == "Mitarbeiter-Shop")
    assert not sh.sperren and pruef.fehler == 0
    lenzing = next(e for e in pruef.einheiten() if e.titel == "Lenzing")
    assert any("Kundenadresse unvollständig (Straße, PLZ, Ort)" in t
               for t in lenzing.warnungen)
    xml = _wex(lenzing, tmp_path)
    sender = xml.split("<Sender>")[1].split("</Sender>")[0]
    for tag in ("Name1", "Street", "PostalCodeCity", "City", "Country"):
        assert f"<{tag} />" in sender
    assert "<DatevNo>10000</DatevNo>" in sender
    # Feste Lieferadresse bleibt davon unberührt
    assert "<Street>Werkplatz 1</Street>" in xml.split("<Delivery>")[1]


def test_platzhalter_in_sender_address_wird_geleert(umgebung):
    _shop(umgebung, "Mitarbeiter-Shop")["sender_address"]["street"] = "BITTE EINTRAGEN — Straße"
    e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "Lenzing")
    assert e.wex_data["street"] == "" and e.wex_data["name1"] == ""
    assert any("(Straße)" in t for t in e.warnungen)


def test_einzel_ohne_lieferanschrift(umgebung, orders, tmp_path):
    o = orders["einzeln"]
    o["shipping"].update({"address_1": "", "postcode": "", "city": ""})
    o["billing"].update({"address_1": "", "postcode": "", "city": ""})
    FakeWoo.orders_by_url["https://shop.example/einzeln/"] = [o]
    e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "402")
    delivery = _wex(e, tmp_path).split("<Delivery>")[1]
    assert "<Street />" in delivery and "<Name1 />" in delivery
    assert any(t.startswith("Lieferanschrift unvollständig") for t in e.warnungen)


def test_einzel_ohne_rechnungsanschrift(umgebung, orders, tmp_path):
    o = orders["einzeln"]
    o["billing"].update({"company": "", "address_1": "", "postcode": "", "city": ""})
    FakeWoo.orders_by_url["https://shop.example/einzeln/"] = [o]
    e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "402")
    xml = _wex(e, tmp_path)
    assert "<Street />" in xml.split("<Sender>")[1].split("</Sender>")[0]
    # Versandadresse der Bestellung bleibt im Delivery-Block
    assert "<Street>Lindenweg 7</Street>" in xml.split("<Delivery>")[1]


def test_kundenadresse_und_lieferung_leer(umgebung, tmp_path):
    """Sammel ohne feste Adresse, Kundenadresse unvollständig, Regel firma:
    die Lieferanschrift hing an der Kundenadresse → ebenfalls leer."""
    agrar = _shop(umgebung, "Agrar-Shop")
    agrar["unknown_delivery"] = "firma"
    agrar["sender_address"] = {"street": "BITTE EINTRAGEN"}
    e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "Bondorf")
    delivery = _wex(e, tmp_path).split("<Delivery>")[1]
    assert "<Name1 />" in delivery and "<Street />" in delivery
    assert "<ModeOfShippment>Bondorf</ModeOfShippment>" in delivery


def test_vollstaendige_adressen_keine_sperre(umgebung):
    pruef = _abrufen(umgebung)
    assert not any(s.sperren for s in pruef.shops)


# --- Prüfregel: Lieferort ohne feste Adresse --------------------------------

def test_lieferort_ohne_adresse_standard_cdh(umgebung, tmp_path):
    e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "Bondorf")
    assert e.wex_data["delivery_leer"] and not e.sperren
    delivery = _wex(e, tmp_path).split("<Delivery>")[1].split("</Delivery>")[0]
    assert "<ModeOfShippment>Bondorf</ModeOfShippment>" in delivery
    for tag in ("Name1", "Name2", "Street", "PostalCodeCity", "City", "Country"):
        assert f"<{tag} />" in delivery
    # Sender bleibt die vollständige Kundenadresse
    assert "<Street />" not in _wex(e, tmp_path).split("<Sender>")[1].split("</Sender>")[0]


def test_lieferort_ohne_adresse_firma(umgebung):
    _shop(umgebung, "Agrar-Shop")["unknown_delivery"] = "firma"
    e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "Bondorf")
    assert e.wex_data["del_street"] == e.wex_data["street"]     # Kundenadresse
    assert not e.wex_data.get("delivery_leer") and not e.sperren


def test_lieferort_ohne_adresse_sperren(umgebung):
    _shop(umgebung, "Agrar-Shop")["unknown_delivery"] = "sperren"
    pruef = _abrufen(umgebung)
    agrar = next(s for s in pruef.shops if s.shop == "Agrar-Shop")
    assert agrar.einheiten[0].sperren and "Bondorf" in agrar.einheiten[0].sperren[0]
    assert "Bondorf" not in [e.titel for e in pruef.einheiten()]
    assert agrar.fehler == 1
    imp = w.importieren(pruef.einheiten())
    assert not any("Bondorf" in n for n in umgebung["cdh"]) and imp.ok == 4


def test_lieferort_ohne_adresse_versandadresse(umgebung, orders):
    _shop(umgebung, "Agrar-Shop")["unknown_delivery"] = "versand"
    e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "Bondorf")
    ship = orders["trenn"][0]["shipping"]
    assert e.wex_data["del_street"] == ship["address_1"]
    assert e.wex_data["del_city"] == ship["city"]
    assert any("Versandadresse der ersten Bestellung (3939)" in t for t in e.warnungen)


def test_lieferort_feste_adresse_hat_vorrang(umgebung):
    _shop(umgebung, "Mitarbeiter-Shop")["unknown_delivery"] = "sperren"
    einheiten = _abrufen(umgebung).einheiten()
    assert {"Seewalchen", "Lenzing"} <= {e.titel for e in einheiten}


def test_unknown_delivery_ungueltig(umgebung, caplog):
    _shop(umgebung, "Agrar-Shop")["unknown_delivery"] = "irgendwas"
    with caplog.at_level(logging.ERROR):
        e = next(e for e in _abrufen(umgebung).einheiten() if e.titel == "Bondorf")
    assert not e.sperren and e.wex_data["delivery_leer"]       # Rückfall: cdh
    assert any("unknown_delivery" in r.getMessage() for r in caplog.records)


# --- Prüfregel: EK fehlt → nur Hinweis --------------------------------------

def test_ek_fehlt_nur_hinweis(umgebung, monkeypatch):
    echt = w.extract_ek_vk

    def ohne_ek(client, item, cache):
        ek, vk = echt(client, item, cache)
        return (None, vk) if item.get("product_id") == 40 else (ek, vk)

    monkeypatch.setattr(w, "extract_ek_vk", ohne_ek)
    pruef = _abrufen(umgebung)
    e = next(e for e in pruef.einheiten() if e.titel == "Bondorf")
    assert any(t.startswith("EK fehlt bei") for t in e.warnungen)
    assert not e.sperren and pruef.fehler == 0


# --- Versandarten -----------------------------------------------------------

def test_versandarten_aus_zonen():
    antworten = {
        "/shipping/zones": [{"id": 0}, {"id": 1}],
        "/shipping/zones/0/methods": [{"title": "Cham", "enabled": True}],
        "/shipping/zones/1/methods": [{"title": "Garbsen"}, {"title": "Cham"},
                                      {"title": "Alt", "enabled": False},
                                      {"method_title": "Abholung"}],
    }
    assert w.versandarten_aus_zonen(antworten.__getitem__) == ["Cham", "Garbsen", "Abholung"]


def test_versandarten_abgleich():
    ab = w.versandarten_abgleich(["Cham", "Seewalchen (Österreich)", "Garbsen "],
                                 ["cham", "Seewalchen", "Lenzing", "Garbsen"])
    assert ab == {"ohne_adresse": ["Seewalchen (Österreich)"],
                  "ohne_versandart": ["Seewalchen", "Lenzing"]}


def test_abrufen_meldet_versandarten(umgebung):
    FakeWoo.zones_by_url["https://shop.example/mitarbeiter/"] = {
        0: [{"title": "Seewalchen"}, {"title": "Cham"}]}
    sh = next(s for s in _abrufen(umgebung).shops if s.shop == "Mitarbeiter-Shop")
    assert sh.versandarten == ["Seewalchen", "Cham"]
    assert "Versandart ohne feste Lieferadresse: Cham" in sh.warnungen
    assert any(t.startswith("Lieferadresse ohne passende Versandart im Shop: Lenzing")
               for t in sh.warnungen)
    assert not sh.sperren


def test_keine_versandarten_ein_hinweis(umgebung):
    sh = next(s for s in _abrufen(umgebung).shops if s.shop == "Mitarbeiter-Shop")
    assert [t for t in sh.warnungen if "Versand" in t] == [
        "Keine aktive Versandart im Shop gefunden — Versandzonen prüfen."]


def test_versandarten_nicht_fuer_einzelshop_ohne_adressen(umgebung):
    _abrufen(umgebung)
    beispiel = [g for g in FakeWoo.gets if g.startswith("/shipping")]
    # Agrar (Sammel) und Mitarbeiter (Sammel) fragen, Beispiel-Shop nicht
    assert len(beispiel) == 2


def test_versandarten_fehler_ist_nur_hinweis(umgebung, monkeypatch):
    def kaputt(self):
        raise RuntimeError("503")
    monkeypatch.setattr(FakeWoo, "get_shipping_methods", kaputt)
    pruef = _abrufen(umgebung)
    agrar = next(s for s in pruef.shops if s.shop == "Agrar-Shop")
    assert any("Versandarten konnten nicht abgerufen werden" in t for t in agrar.warnungen)
    assert len(pruef.einheiten()) == 4


# --- Excel: Summenblatt und Übersicht ---------------------------------------

def test_summen_zeilen_je_ort_und_gesamt(orders):
    alle = orders["mitarbeitershop"]
    head, rows = w.summen_zeilen(alle, by="ort")
    assert head[0] == "Lieferort" and head[-1] == "Anzahl"
    assert sum(r[-1] for r in rows) == sum(i["quantity"] for o in alle for i in o["line_items"])
    assert {r[0] for r in rows} == {"Seewalchen", "Lenzing"}

    head, rows = w.summen_zeilen(alle, by="gesamt")
    assert head[0] == "Artikelnummer"
    assert len(rows) == len({(i["sku"], w._variant_text_labeled(i))
                             for o in alle for i in o["line_items"]})


def test_summen_zeilen_ohne_veredelungen(orders):
    trenn = orders["trenn"]
    _, mit = w.summen_zeilen(trenn, by="gesamt")
    _, ohne = w.summen_zeilen(trenn, by="gesamt", mit_veredelungen=False)
    assert any(w._is_veredelung(r[0]) for r in mit)
    assert not any(w._is_veredelung(r[0]) for r in ohne)


def test_import_excel_mit_summenblatt(orders, client, tmp_path):
    out = tmp_path / "t.xlsx"
    w.write_excel_export(out, orders["trenn"], {"datev_no": 1, "excel_summary": True,
                                                "excel_summary_by": "gesamt"}, client, {})
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Bestellungen", "Summe"]
    assert [c.value for c in wb["Summe"][1]] == ["Artikelnummer", "Artikelname",
                                                 "Artikeltext 2", "Anzahl"]


def test_import_excel_ohne_summenblatt(orders, client, tmp_path):
    out = tmp_path / "t.xlsx"
    w.write_excel_export(out, orders["trenn"], {"datev_no": 1}, client, {})
    assert openpyxl.load_workbook(out).sheetnames == ["Bestellungen"]


def test_excel_uebersicht_ohne_import(umgebung):
    _shop(umgebung, "Mitarbeiter-Shop")["excel_summary"] = True
    pruef = _abrufen(umgebung)
    out = umgebung["tmp"] / "uebersicht.xlsx"
    w.write_excel_uebersicht(out, pruef.shops)

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Beispiel-Shop", "Agrar-Shop", "Mitarbeiter-Shop",
                             "Summe Mitarbeiter-Shop"]
    assert wb["Agrar-Shop"].max_row == 1 + 4          # Kopf + 4 Positionen
    # Nichts importiert, nichts markiert
    assert FakeWoo.puts == [] and umgebung["cdh"] == []
    assert not w.EXPORTED_LOG_PATH.exists()
    assert not (umgebung["tmp"] / "wex").exists()


def test_blattname():
    assert w._blattname("A/B:C*D?[E]") == "A-B-C-D--E-"
    assert len(w._blattname("x" * 40)) == 31


# --- Importsperre über Rechner ----------------------------------------------

@pytest.fixture
def lock(tmp_path, monkeypatch):
    p = tmp_path / "running.lock"
    monkeypatch.setattr(w, "LOCK_PATH", p)
    monkeypatch.setenv("COMPUTERNAME", "PC-TEST")
    monkeypatch.setenv("USERNAME", "tester")
    yield p
    w.release_lock()


def _fremd(p, alter_sek=0):
    p.write_text(json.dumps({"rechner": "PC-LAGER", "benutzer": "m.mueller",
                             "pid": 1, "seit": "2026-10-01T09:14:00"}), encoding="utf-8")
    t = time.time() - alter_sek
    os.utime(p, (t, t))


def test_lock_mit_rechner_und_zeit(lock):
    assert w.acquire_lock()
    info = json.loads(lock.read_text(encoding="utf-8"))
    assert info["rechner"] == "PC-TEST" and info["benutzer"] == "tester"
    assert "seit" in info and info["pid"] == os.getpid()
    w.release_lock()
    assert not lock.exists()


def test_anderer_rechner_sieht_import_laeuft(lock, capsys):
    _fremd(lock)
    assert w.lock_text(w.lock_info()) == "Import läuft an PC-LAGER (m.mueller) seit 09:14"
    assert not w.acquire_lock()
    assert "Import läuft an PC-LAGER (m.mueller) seit 09:14" in capsys.readouterr().out
    assert lock.exists()


def test_verwaiste_sperre_wird_uebernommen(lock):
    _fremd(lock, alter_sek=w.LOCK_STALE_MINUTES * 60 + 5)
    assert w.lock_info() is None
    assert w.acquire_lock()
    assert json.loads(lock.read_text(encoding="utf-8"))["rechner"] == "PC-TEST"


def test_heartbeat_haelt_sperre_frisch(lock, monkeypatch):
    monkeypatch.setattr(w, "LOCK_HEARTBEAT_SECONDS", 0.05)
    assert w.acquire_lock()
    alt = time.time() - 3600
    os.utime(lock, (alt, alt))
    time.sleep(0.3)
    assert time.time() - lock.stat().st_mtime < 5
    assert w.lock_info() is not None


def test_release_loescht_fremde_sperre_nicht(lock):
    assert w.acquire_lock()
    _fremd(lock)                     # anderer Rechner hat übernommen
    w.release_lock()
    assert lock.exists()


def test_alte_freitext_sperre(lock):
    lock.write_text("PC-ALT\\user PID 5 @ 2026-09-01T08:00:00", encoding="utf-8")
    info = w.lock_info()
    assert info["rechner"].startswith("PC-ALT")
    assert w.lock_text(info).startswith("Import läuft an PC-ALT")


# --- Diagnose als Funktion --------------------------------------------------

def _diag_api(orders, dims=None, status=200, zones=None):
    dims = dims or {}
    zones = zones if zones is not None else {0: [{"title": "Standardversand"}]}

    def abruf(path, params=None):
        if status != 200:
            return status, {}
        if path == "/orders":
            return 200, orders
        if path == "/shipping/zones":
            return 200, [{"id": z} for z in zones]
        if path.startswith("/shipping/zones/"):
            return 200, zones[int(path.split("/")[3])]
        parts = path.strip("/").split("/")
        key = (int(parts[1]), int(parts[3]) if len(parts) > 3 else 0)
        ek, vk = dims.get(key, ("1.00", "2.00"))
        return 200, {"dimensions": {"length": ek, "width": vk}}
    return abruf


def test_diagnose_alles_ok(orders):
    punkte = d.diagnose({"name": "X", "url": "https://x/"}, _diag_api([orders["einzeln"]]))
    assert [p.titel for p in punkte] == ["Verbindung und Zugang", "Bestellungen lesbar",
                                         "Preise gepflegt", "Varianten erkannt", "Versandarten"]
    assert [p.stufe for p in punkte] == ["ok"] * 5
    assert punkte[1].text == "1 offene Bestellung gefunden"
    assert punkte[4].text == "1 gefunden: Standardversand"


def test_diagnose_401_bricht_ab():
    punkte = d.diagnose({"url": "https://x/"}, _diag_api([], status=401))
    assert len(punkte) == 1 and punkte[0].stufe == "fehler"
    assert "Sub-Shop" in punkte[0].text


def test_diagnose_ek_fehlt(orders):
    punkte = d.diagnose({"url": "https://x/"},
                        _diag_api(orders["trenn"], dims={(40, 0): ("", "7.65")}))
    preise = punkte[2]
    assert preise.stufe == "warn"
    assert preise.text.startswith("EK fehlt bei 1 von 2 Artikeln")
    assert preise.details == [orders["trenn"][0]["line_items"][1]["sku"]]


def test_diagnose_ohne_bestellungen():
    punkte = d.diagnose({"url": "https://x/"}, _diag_api([]))
    assert [p.stufe for p in punkte] == ["ok", "warn", "warn", "warn", "ok"]


def test_diagnose_versandarten_abgleich(orders, monkeypatch):
    from conftest import FIXTURES
    monkeypatch.setattr(w, "DELIVERY_ADDRESSES_PATH", FIXTURES / "lieferadressen_test.yaml")
    punkte = d.diagnose({"name": "Mitarbeiter-Shop", "url": "https://x/"},
                        _diag_api(orders["mitarbeitershop"],
                                  zones={1: [{"title": "Seewalchen (Österreich)"}]}))
    va = punkte[4]
    assert va.stufe == "warn"
    assert "Adresse ohne Versandart: Seewalchen, Lenzing" in va.details


def test_diagnose_ohne_netz_schreibt_nichts(orders):
    aufrufe = []
    api = _diag_api([orders["einzeln"]])
    d.diagnose({"url": "https://x/"}, lambda p, params=None: aufrufe.append(p) or api(p, params))
    assert all(p.startswith(("/orders", "/products", "/shipping")) for p in aufrufe)
